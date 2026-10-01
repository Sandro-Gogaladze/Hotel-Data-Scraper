#!/usr/bin/env python3
"""
Retry failed hotel detail extractions from an existing Booking_Data workbook.

The normal exporter keeps a boolean "Execution Error" column in the final Excel
file. This utility uses that flag to retry only failed rows, then writes a
filled copy of the workbook with successful retry values patched into place.
"""

import argparse
import os
import sys
from datetime import datetime
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from playwright.sync_api import sync_playwright

project_root = os.path.abspath(os.path.dirname(__file__))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

from config import EXCEL_ERROR_COLOR, HEADLESS, OUTPUT_DIR
from scraper.hotel_page.detailed_extractor import extract_detailed_info
from utils.anti_detection import add_random_delay, get_browser_launch_options
from utils.logger import log_message


DETAIL_COLUMNS = [
    "Min Room Price",
    "Min Room Price Free Cancellation",
    "Min Room Price Non-refundable",
    "Min Room Price Breakfast Included",
    "Min 2 Person Price",
    "Min 2 Person Price Free Cancellation",
    "Min 2 Person Price Non-refundable",
    "Min 2 Person Price Breakfast Included",
]


def latest_workbook() -> str:
    candidates = [
        os.path.join(OUTPUT_DIR, name)
        for name in os.listdir(OUTPUT_DIR)
        if name.lower().endswith(".xlsx") and not name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(f"No .xlsx files found in {OUTPUT_DIR}")
    return max(candidates, key=os.path.getmtime)


def default_output_path(input_path: str) -> str:
    base, ext = os.path.splitext(input_path)
    timestamp = datetime.now().strftime("%H-%M-%S")
    return f"{base}_retry_filled_{timestamp}{ext}"


def header_map(ws: Any) -> Dict[str, int]:
    headers = {}
    for cell in ws[1]:
        if cell.value:
            headers[str(cell.value)] = cell.column
    return headers


def failed_rows(ws: Any, headers: Dict[str, int]) -> List[int]:
    execution_error_col = headers["Execution Error"]
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        value = ws.cell(row_idx, execution_error_col).value
        if value is True or value == 1 or str(value).strip().upper() == "TRUE":
            rows.append(row_idx)
    return rows


def hotel_from_row(ws: Any, headers: Dict[str, int], row_idx: int) -> Dict[str, Any]:
    link_cell = ws.cell(row_idx, headers["Hotel Link"])
    hotel_link = link_cell.hyperlink.target if link_cell.hyperlink else link_cell.value
    return {
        "Index": row_idx - 1,
        "Hotel Name": ws.cell(row_idx, headers["Hotel Name"]).value,
        "Address": ws.cell(row_idx, headers["Address"]).value,
        "Price": ws.cell(row_idx, headers["Price"]).value,
        "Review Score": ws.cell(row_idx, headers["Review Score"]).value,
        "Review Text": ws.cell(row_idx, headers["Review Text"]).value,
        "Number of Reviews": ws.cell(row_idx, headers["Number of Reviews"]).value,
        "Stars": ws.cell(row_idx, headers["Stars"]).value,
        "Hotel Link": hotel_link,
    }


def is_successful_retry(result: Dict[str, Any]) -> bool:
    return (
        not str(result.get("Extraction Error", "")).strip()
        and result.get("Min Room Price") not in (None, "", "N/A")
    )


def patch_row(ws: Any, headers: Dict[str, int], row_idx: int, result: Dict[str, Any]) -> None:
    for col_name in DETAIL_COLUMNS:
        ws.cell(row_idx, headers[col_name], result.get(col_name, "N/A"))

    execution_error_cell = ws.cell(row_idx, headers["Execution Error"])
    execution_error_cell.value = ""
    execution_error_cell.fill = PatternFill(fill_type=None)

    for col_name in DETAIL_COLUMNS:
        cell = ws.cell(row_idx, headers[col_name])
        if cell.value == "N/A":
            cell.fill = PatternFill(
                start_color=EXCEL_ERROR_COLOR,
                end_color=EXCEL_ERROR_COLOR,
                fill_type="solid",
            )


def retry_hotels(input_path: str, output_path: str, max_retries: int, headless: bool) -> Tuple[int, int, List[str]]:
    wb = load_workbook(input_path)
    ws = wb["Hotel Data"]
    headers = header_map(ws)
    required = {"Hotel Name", "Hotel Link", "Execution Error", *DETAIL_COLUMNS}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(missing)}")

    rows = failed_rows(ws, headers)
    if not rows:
        wb.save(output_path)
        return 0, 0, []

    failures: List[str] = []
    success_count = 0

    with sync_playwright() as p:
        launch_options = get_browser_launch_options(headless=headless, worker_id=0)
        browser = p.chromium.launch(**launch_options)
        page = browser.new_page()
        page.set_default_timeout(60000)

        try:
            for position, row_idx in enumerate(rows, start=1):
                hotel = hotel_from_row(ws, headers, row_idx)
                hotel_name = hotel.get("Hotel Name", f"row {row_idx}")
                log_message(
                    f"[Retry] {position}/{len(rows)} row {row_idx}: {hotel_name}",
                    "info",
                )
                add_random_delay(1.0, 2.0)
                result = extract_detailed_info(
                    browser,
                    hotel,
                    page=page,
                    max_retries=max_retries,
                )

                if is_successful_retry(result):
                    patch_row(ws, headers, row_idx, result)
                    success_count += 1
                    log_message(f"[Retry] Filled row {row_idx}: {hotel_name}", "info")
                else:
                    error = result.get("Extraction Error", "Unknown retry failure")
                    failures.append(f"row {row_idx} {hotel_name}: {error}")
                    log_message(f"[Retry] Still failed row {row_idx}: {hotel_name} - {error}", "warning")
        finally:
            page.close()
            browser.close()

    wb.save(output_path)
    return len(rows), success_count, failures


def main() -> None:
    parser = argparse.ArgumentParser(description="Retry failed hotel rows in a Booking scraper Excel workbook.")
    parser.add_argument("workbook", nargs="?", default=None, help="Workbook to patch. Defaults to latest in Booking_Data.")
    parser.add_argument("--output", default=None, help="Output workbook path. Defaults to *_retry_filled_<time>.xlsx.")
    parser.add_argument("--max-retries", type=int, default=4, help="Retries per failed hotel.")
    parser.add_argument("--headed", action="store_true",
                        help="Run Chromium with a visible window (the default unless HEADLESS=true).")
    args = parser.parse_args()

    input_path = args.workbook or latest_workbook()
    output_path = args.output or default_output_path(input_path)

    total, filled, failures = retry_hotels(
        input_path=input_path,
        output_path=output_path,
        max_retries=args.max_retries,
        headless=HEADLESS and not args.headed,
    )

    print(f"Input workbook: {input_path}")
    print(f"Output workbook: {output_path}")
    print(f"Failed rows retried: {total}")
    print(f"Rows filled: {filled}")
    print(f"Rows still failing: {len(failures)}")
    for failure in failures:
        print(f"- {failure}")


if __name__ == "__main__":
    main()
